"""Synthetic Office create/edit/reopen/render qualification; no user files or installs.

Optional --fixtures accepts only a directory of the synthetic fixtures exported by this script.
--export-fixtures writes a NEW task-specific directory; it never replaces an existing directory.
"""
import argparse
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

sys.dont_write_bytecode = True
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--assets", type=Path, default=Path(__file__).resolve().parent.parent / "resources/skills")
parser.add_argument("--fixtures", type=Path)
parser.add_argument("--export-fixtures", type=Path)
args = parser.parse_args()
ASSETS = args.assets.resolve()


def create_fixtures(root):
    from docx import Document
    from openpyxl import Workbook
    doc = Document()
    doc.add_heading("Office qualification", 0)
    doc.add_paragraph("Keep this café & XML <sample> paragraph.")
    doc.add_paragraph("Review target paragraph 42.")
    doc.save(root / "seed.docx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fixture"
    sheet["A1"] = 19
    sheet["A2"] = 23
    sheet["A3"] = "=SUM(A1:A2)"
    sheet["B1"] = "Office fixture"
    workbook.save(root / "seed.xlsx")
    if importlib.util.find_spec("pptx"):
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Office qualification"
        slide.placeholders[1].text = "Synthetic slide 42"
        presentation.save(root / "seed.pptx")


if args.export_fixtures:
    args.export_fixtures.mkdir(parents=True, exist_ok=False)
    create_fixtures(args.export_fixtures)
    print(str(args.export_fixtures.resolve()))
    sys.exit(0)


class OfficeWorkflows(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="sidekick-office-workflows-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        if args.fixtures:
            for name in ("seed.docx", "seed.xlsx", "seed.pptx"):
                source = args.fixtures / name
                if source.exists():
                    shutil.copy2(source, self.root / name)
        else:
            create_fixtures(self.root)

    def helper(self, relative, *arguments, expected=0, timeout=45):
        result = subprocess.run([sys.executable, "-B", str(ASSETS / relative), *map(str, arguments)],
                                cwd=self.root, capture_output=True, text=True, timeout=timeout)
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return result.stdout

    def unpack(self, name="seed.docx"):
        target = self.root / "unpacked"
        self.helper("office/unpack.py", self.root / name, target, "--merge-runs", "false", "--simplify-redlines", "false")
        return target

    def snapshot(self, root):
        return {p.relative_to(root).as_posix(): p.read_bytes() for p in root.rglob("*") if p.is_file()}

    def test_comments_create_extend_reject_and_reopen(self):
        from docx import Document
        from lxml import etree as ET
        source = self.root / "seed.docx"
        original = source.read_bytes()
        unpacked = self.unpack()
        before = self.snapshot(unpacked)
        self.helper("docx/comment.py", unpacked, "auto", 'Review <XML> & café "42"', "--paragraph", 2, "--author", 'Test & "Author"')
        first = self.snapshot(unpacked)
        changed = {name for name in first if before.get(name) != first[name]}
        self.assertEqual(changed, {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"})
        self.helper("docx/comment.py", unpacked, "auto", "Second comment", "--paragraph", 1)
        complete = self.snapshot(unpacked)
        for arguments in [(0, "Duplicate", "--paragraph", 1), ("auto", "Reply", "--paragraph", 1, "--parent", 0), ("auto", "Missing target", "--paragraph", 999)]:
            self.helper("docx/comment.py", unpacked, *arguments, expected=1)
            self.assertEqual(self.snapshot(unpacked), complete)
        report = json.loads(self.helper("office/validate.py", unpacked))
        self.assertTrue(report["valid"])
        self.assertFalse(report["xsd_validation"])
        result = self.root / "commented.docx"
        self.helper("office/pack.py", unpacked, result)
        self.assertEqual([p.text for p in Document(source).paragraphs], [p.text for p in Document(result).paragraphs])
        with zipfile.ZipFile(result) as archive:
            comments = ET.fromstring(archive.read("word/comments.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            self.assertEqual(comments.xpath("w:comment/@w:id", namespaces=ns), ["0", "1"])
            self.assertEqual(comments.xpath("w:comment[1]//w:t/text()", namespaces=ns), ['Review <XML> & café "42"'])
            self.assertEqual(comments[0].get("{" + ns["w"] + "}author"), 'Test & "Author"')
        self.assertEqual(source.read_bytes(), original)

    def test_structural_validator_rejects_broken_relationship_and_comment_anchor(self):
        from lxml import etree as ET
        unpacked = self.unpack()
        rels = unpacked / "word/_rels/document.xml.rels"
        tree = ET.fromstring(rels.read_bytes())
        tree[0].set("Target", "missing-part.xml")
        rels.write_bytes(ET.tostring(tree))
        report = json.loads(self.helper("office/validate.py", unpacked, expected=1))
        self.assertTrue(any("Missing relationship target" in message for message in report["errors"]))
        output = self.root / "rejected.docx"
        self.helper("office/pack.py", unpacked, output, expected=1)
        self.assertFalse(output.exists())
        # Restore the relationship, then reject an independently broken comment range.
        original = zipfile.ZipFile(self.root / "seed.docx")
        self.addCleanup(original.close)
        rels.write_bytes(original.read("word/_rels/document.xml.rels"))
        self.helper("docx/comment.py", unpacked, "auto", "Fixture", "--paragraph", 1)
        document_path = unpacked / "word/document.xml"
        document = ET.fromstring(document_path.read_bytes())
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        end = next(document.iter(f"{{{ns}}}commentRangeEnd"))
        end.getparent().remove(end)
        document_path.write_bytes(ET.tostring(document))
        report = json.loads(self.helper("office/validate.py", unpacked, expected=1))
        self.assertIn("Unpaired comment range markers", report["errors"])

    def test_structural_validator_rejects_zip_traversal_and_xml_entities(self):
        malicious = self.root / "bad.docx"
        with zipfile.ZipFile(malicious, "w") as archive:
            archive.writestr("../outside.xml", "fixture")
        report = json.loads(self.helper("office/validate.py", malicious, expected=1))
        self.assertIn("Unsafe package member", report["errors"][0])
        rejected = self.root / "rejected-unpack"
        self.helper("office/unpack.py", malicious, rejected, expected=1)
        self.assertFalse(rejected.exists())
        unpacked = self.unpack()
        (unpacked / "word/document.xml").write_text('<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>')
        report = json.loads(self.helper("office/validate.py", unpacked, expected=1))
        self.assertTrue(any("DTD declarations" in message for message in report["errors"]))

    def test_structural_validator_checks_xlsx_and_pptx_fixtures(self):
        for name in ("seed.xlsx", "seed.pptx"):
            path = self.root / name
            if not path.exists():
                self.skipTest("Synthetic PPTX fixture unavailable; no python-pptx installation attempted")
            report = json.loads(self.helper("office/validate.py", path))
            self.assertEqual(report["scope"], "opc-ooxml-structural")
        rejected = json.loads(self.helper("office/validate.py", self.root / "seed.docx", "--xsd", expected=1))
        self.assertEqual(rejected["scope"], "xsd")

    def test_libreoffice_recalculates_exact_formula_without_touching_source(self):
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice unavailable; success path not qualified")
        from openpyxl import load_workbook
        source = self.root / "seed.xlsx"
        before = source.read_bytes()
        report = json.loads(self.helper("xlsx/recalc.py", source, 30))
        self.assertEqual(report["status"], "success")
        self.assertEqual(report["total_formulas"], 1)
        output = Path(report["output"])
        workbook = load_workbook(output, data_only=True)
        self.assertEqual(workbook["Fixture"]["A3"].value, 42)
        workbook.close()
        self.assertEqual(source.read_bytes(), before)
        self.helper("xlsx/recalc.py", source, 30, expected=1)

    def test_libreoffice_accepts_tracked_changes_and_reopens(self):
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice unavailable; tracked-change success path not qualified")
        from lxml import etree as ET
        from docx import Document
        unpacked = self.unpack()
        path = unpacked / "word/document.xml"
        tree = ET.fromstring(path.read_bytes())
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        paragraph = tree.find(f"{{{ns}}}body").findall(f"{{{ns}}}p")[2]
        for tag, text_tag, text in [("del", "delText", "Remove this old phrase"), ("ins", "t", "Accepted new phrase")]:
            change = ET.SubElement(paragraph, f"{{{ns}}}{tag}", {f"{{{ns}}}id": "100" if tag == "del" else "101", f"{{{ns}}}author": "Fixture", f"{{{ns}}}date": "2026-09-05T00:00:00Z"})
            run = ET.SubElement(change, f"{{{ns}}}r")
            ET.SubElement(run, f"{{{ns}}}{text_tag}").text = text
        path.write_bytes(ET.tostring(tree))
        tracked = self.root / "tracked.docx"
        self.helper("office/pack.py", unpacked, tracked)
        before = tracked.read_bytes()
        accepted = self.root / "accepted.docx"
        self.helper("docx/accept_changes.py", tracked, accepted, timeout=50)
        text = "\n".join(p.text for p in Document(accepted).paragraphs)
        self.assertIn("Accepted new phrase", text)
        self.assertNotIn("Remove this old phrase", text)
        self.assertEqual(tracked.read_bytes(), before)

    def test_libreoffice_renders_commented_docx_and_pptx_to_readable_pdf(self):
        if not shutil.which("soffice") or not shutil.which("pdftotext"):
            self.skipTest("LibreOffice/Poppler unavailable; rendering not qualified")
        unpacked = self.unpack()
        self.helper("docx/comment.py", unpacked, "auto", "Render-check comment", "--paragraph", 2)
        commented = self.root / "commented.docx"
        self.helper("office/pack.py", unpacked, commented)
        for name, expected in [("commented.docx", "Review target paragraph 42"), ("seed.pptx", "Synthetic slide 42")]:
            source = self.root / name
            self.assertTrue(source.exists(), "Synthetic rendering fixture missing")
            destination = self.root / (source.stem + "-rendered")
            result = json.loads(self.helper("office/render.py", source, destination, "--images", timeout=90))
            self.assertEqual(result["status"], "success")
            pdf = Path(result["pdf"])
            self.assertTrue(pdf.is_file())
            self.assertGreater(len(result["images"]), 0)
            self.assertTrue(all(Path(path).read_bytes().startswith(b"\x89PNG") for path in result["images"]))
            text = subprocess.check_output(["pdftotext", str(pdf), "-"], text=True, timeout=10)
            self.assertIn(expected, text)


if __name__ == "__main__":
    unittest.main(argv=[sys.argv[0]], verbosity=2)
