"""Recalculate into a NEW workbook using an isolated LibreOffice profile.

Never installs software, changes the user's LibreOffice profile, or overwrites input.
"""

import argparse
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile


def recalc(filename, timeout=30, output=None):
    source = Path(filename).resolve()
    destination = Path(output).resolve() if output else source.with_name(source.stem + "-recalculated.xlsx")
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        return {"error": "Input must be an existing .xlsx file"}
    if destination == source or destination.exists():
        return {"error": "Output must be a new file; source and existing outputs are preserved"}
    if destination.suffix.lower() != ".xlsx":
        return {"error": "Output must have the .xlsx extension"}
    executable = shutil.which("soffice")
    if not executable:
        return {"error": "Missing dependency: LibreOffice (soffice on PATH). Nothing installed or modified."}
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"error": "Missing dependency: openpyxl. Nothing installed or modified."}
    try:
        with tempfile.TemporaryDirectory(prefix="sidekick-recalc-") as temporary:
            root = Path(temporary)
            profile = (root / "profile").as_uri()
            converted_dir = root / "converted"
            converted_dir.mkdir()
            result = subprocess.run(
                [executable, f"-env:UserInstallation={profile}", "--headless", "--norestore",
                 "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(converted_dir), str(source)],
                capture_output=True, text=True, timeout=max(1, timeout), check=False,
            )
            converted = converted_dir / source.name
            if result.returncode != 0 or not converted.is_file():
                return {"error": "LibreOffice conversion failed; no output published"}
            values = load_workbook(converted, data_only=True)
            formulas = load_workbook(converted, data_only=False)
            try:
                original = load_workbook(source, data_only=False)
                try:
                    for sheet in original:
                        for row in sheet:
                            for cell in row:
                                if cell.data_type == "f" and (sheet.title not in formulas.sheetnames or formulas[sheet.title][cell.coordinate].data_type != "f"):
                                    return {"error": "Conversion lost an original formula; no output published"}
                finally:
                    original.close()
                formula_count = 0
                missing_cache = []
                errors = []
                for sheet in formulas:
                    for row in sheet:
                        for cell in row:
                            value = values[sheet.title][cell.coordinate]
                            if cell.data_type == "f":
                                formula_count += 1
                                if value.value is None:
                                    missing_cache.append(f"{sheet.title}!{cell.coordinate}")
                            if value.data_type == "e":
                                errors.append(f"{sheet.title}!{cell.coordinate}: {value.value}")
                if missing_cache:
                    return {"error": "Formula caches remain missing; no output published", "cells": missing_cache[:20]}
                # Exclusive creation protects an output created since the initial check.
                with destination.open("xb") as target, converted.open("rb") as content:
                    shutil.copyfileobj(content, target)
                return {"status": "errors_found" if errors else "success", "output": str(destination),
                        "total_formulas": formula_count, "total_errors": len(errors), "errors": errors[:20],
                        "note": "LibreOffice conversion; verify formatting and unsupported Excel features separately."}
            finally:
                values.close()
                formulas.close()
    except subprocess.TimeoutExpired:
        return {"error": "LibreOffice timed out; recalculation is not verified and no output was published"}
    except Exception as error:
        return {"error": str(error)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("filename")
    parser.add_argument("timeout", nargs="?", type=int, default=30)
    parser.add_argument("--output", help="New .xlsx destination; defaults to NAME-recalculated.xlsx")
    args = parser.parse_args()
    report = recalc(args.filename, args.timeout, args.output)
    print(json.dumps(report, indent=2))
    sys.exit(0 if report.get("status") == "success" else 1)
