"""Positive/negative oracle checks; disposable synthetic documents only."""
import importlib.util
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.dont_write_bytecode = True
from fixture import create, verify

HELPERS = Path(__file__).resolve().parents[5] / "resources" / "skills"


class OfficeOracleTests(unittest.TestCase):
    def test_docx(self):
        with tempfile.TemporaryDirectory(prefix="sidekick-office-oracle-") as directory:
            root = Path(directory)
            create("docx-comment", root)
            commands = [
                [HELPERS / "office/unpack.py", root / "source.docx", root / "unpacked"],
                [HELPERS / "docx/comment.py", root / "unpacked", "auto", "Confirm the shipping date.", "--paragraph", "2", "--author", "Office QA"],
                [HELPERS / "office/pack.py", root / "unpacked", root / "reviewed.docx"],
            ]
            for command in commands:
                subprocess.run([sys.executable, "-B", *map(str, command)], check=True, capture_output=True)
            self.assertTrue(verify("docx-comment", root)["passed"])
            (root / "reviewed.docx").write_bytes((root / "source.docx").read_bytes())
            with self.assertRaises(KeyError):
                verify("docx-comment", root)

    def test_xlsx(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(prefix="sidekick-office-oracle-") as directory:
            root = Path(directory)
            create("xlsx-edit", root)
            workbook = load_workbook(root / "source.xlsx")
            workbook["Orders"]["B2"] = 7
            workbook.save(root / "updated.xlsx")
            self.assertTrue(verify("xlsx-edit", root)["passed"])
            workbook["Archive"]["A1"] = "wrong"
            workbook.save(root / "updated.xlsx")
            workbook.close()
            self.assertFalse(verify("xlsx-edit", root)["passed"])

    def test_pptx(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory(prefix="sidekick-office-oracle-") as directory:
            root = Path(directory)
            create("pptx-edit", root)
            deck = Presentation(root / "source.pptx")
            deck.slides[1].shapes.title.text = "Shipping checklist — reviewed"
            deck.save(root / "revised.pptx")
            self.assertTrue(verify("pptx-edit", root)["passed"])
            deck.slides[0].shapes.title.text = "wrong"
            deck.save(root / "revised.pptx")
            self.assertFalse(verify("pptx-edit", root)["passed"])

    def test_xlsx_preflight_does_not_require_legacy_xls_reader(self):
        spec = importlib.util.spec_from_file_location("preflight", HELPERS / "preflight.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with patch.object(module.importlib.util, "find_spec", side_effect=lambda name: object() if name == "openpyxl" else None):
            self.assertEqual(module.inspect("xlsx", Path.cwd())["status"], "available")
            self.assertEqual(module.inspect("xls-read", Path.cwd())["status"], "missing_dependencies")


if __name__ == "__main__":
    unittest.main()
