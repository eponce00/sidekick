"""External synthetic fixture builder/verifier for model-driven Office tasks.

Never copied into the model workspace. JSON output contains only checks, not document text.
"""
import argparse
from copy import copy
import json
from pathlib import Path
import sys
import zipfile

sys.dont_write_bytecode = True


def create(scenario, root):
    root = Path(root)
    if scenario == "docx-comment":
        from docx import Document
        doc = Document()
        doc.add_heading("Delivery memo", 0)
        doc.add_paragraph("Reference code: OFFICE-42. Preserve this paragraph.")
        doc.add_paragraph("Please confirm delivery on Monday.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Count"
        table.cell(1, 0).text, table.cell(1, 1).text = "Boxes", "3"
        doc.sections[0].header.paragraphs[0].text = "Synthetic header: keep unchanged"
        doc.save(root / "source.docx")
    elif scenario == "xlsx-edit":
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Orders"
        ws.append(["Item", "Quantity", "Unit price", "Total"])
        ws.append(["Packing boxes", 3, 4.5, "=B2*C2"])
        ws.append(["Labels", 10, 0.25, "=B3*C3"])
        ws["B2"].fill = PatternFill("solid", fgColor="DCE6F1")
        ws["D2"].number_format = '"$"#,##0.00'
        wb.create_sheet("Archive")["A1"] = "Keep this archived record: OFFICE-42"
        wb["Archive"].sheet_state = "hidden"
        wb.save(root / "source.xlsx")
    elif scenario == "pptx-edit":
        from pptx import Presentation
        deck = Presentation()
        first = deck.slides.add_slide(deck.slide_layouts[1])
        first.shapes.title.text = "Planning brief"
        first.placeholders[1].text = "Preserve the first slide: OFFICE-42"
        second = deck.slides.add_slide(deck.slide_layouts[1])
        second.shapes.title.text = "Shipping checklist"
        second.placeholders[1].text = "Pack\nLabel\nDispatch"
        second.notes_slide.notes_text_frame.text = "Keep these speaker notes: OFFICE-42"
        deck.save(root / "source.pptx")
    else:
        raise ValueError("Unknown scenario")


def verify(scenario, root):
    root = Path(root)
    checks = []

    def check(name, condition):
        checks.append({"name": name, "passed": bool(condition)})

    if scenario == "docx-comment":
        from docx import Document
        from lxml import etree as ET
        source, output = Document(root / "source.docx"), Document(root / "reviewed.docx")
        check("paragraphs_preserved", [p.text for p in source.paragraphs] == [p.text for p in output.paragraphs])
        check("tables_preserved", [[[c.text for c in r.cells] for r in t.rows] for t in source.tables] == [[[c.text for c in r.cells] for r in t.rows] for t in output.tables])
        check("header_preserved", source.sections[0].header.paragraphs[0].text == output.sections[0].header.paragraphs[0].text)
        with zipfile.ZipFile(root / "reviewed.docx") as archive:
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            comments = ET.fromstring(archive.read("word/comments.xml"))
            comment = comments.findall("w:comment", ns)
            check("one_exact_comment", len(comment) == 1 and "".join(comment[0].itertext()) == "Confirm the shipping date.")
            check("comment_author", len(comment) == 1 and comment[0].get("{" + ns["w"] + "}author") == "Office QA")
            document = ET.fromstring(archive.read("word/document.xml"))
            marked = document.xpath("//w:p[w:commentRangeStart]", namespaces=ns)
            check("correct_anchor", len(marked) == 1 and "".join(marked[0].xpath(".//w:t/text()", namespaces=ns)) == "Please confirm delivery on Monday.")
            cid = comment[0].get("{" + ns["w"] + "}id") if comment else None
            for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
                check("matching_" + tag, document.xpath("//w:" + tag + "/@w:id", namespaces=ns) == [cid])
    elif scenario == "xlsx-edit":
        from openpyxl import load_workbook
        source, output = load_workbook(root / "source.xlsx"), load_workbook(root / "updated.xlsx")
        try:
            check("sheet_names_preserved", source.sheetnames == output.sheetnames)
            check("quantity_updated", output["Orders"]["B2"].value == 7)
            unchanged = True
            formatting = True
            for sheet in source:
                for row in sheet:
                    for cell in row:
                        other = output[sheet.title][cell.coordinate]
                        formatting &= all(copy(getattr(cell, prop)) == copy(getattr(other, prop))
                                          for prop in ("font", "fill", "border", "alignment", "protection", "number_format"))
                        if sheet.title == "Orders" and cell.coordinate == "B2":
                            continue
                        unchanged &= output[sheet.title][cell.coordinate].value == cell.value
            check("other_cells_and_formulas_preserved", unchanged)
            check("all_fixture_cell_formatting_preserved", formatting)
            check("input_fill_preserved", copy(source["Orders"]["B2"].fill) == copy(output["Orders"]["B2"].fill))
            check("currency_format_preserved", source["Orders"]["D2"].number_format == output["Orders"]["D2"].number_format)
            check("archive_hidden", output["Archive"].sheet_state == "hidden")
        finally:
            source.close()
            output.close()
    elif scenario == "pptx-edit":
        from pptx import Presentation
        source, output = Presentation(root / "source.pptx"), Presentation(root / "revised.pptx")
        check("slide_count", len(output.slides) == len(source.slides) == 2)
        check("page_size_preserved", (source.slide_width, source.slide_height) == (output.slide_width, output.slide_height))
        check("title_updated", output.slides[1].shapes.title.text == "Shipping checklist — reviewed")
        check("first_slide_preserved", [s.text for s in source.slides[0].shapes if s.has_text_frame] == [s.text for s in output.slides[0].shapes if s.has_text_frame])
        check("body_preserved", source.slides[1].placeholders[1].text == output.slides[1].placeholders[1].text)
        check("speaker_notes_preserved", source.slides[1].notes_slide.notes_text_frame.text == output.slides[1].notes_slide.notes_text_frame.text)
    else:
        raise ValueError("Unknown scenario")
    return {"passed": all(c["passed"] for c in checks), "checks": checks}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("operation", choices=("create", "verify"))
    parser.add_argument("scenario", choices=("docx-comment", "xlsx-edit", "pptx-edit"))
    parser.add_argument("workspace", type=Path)
    args = parser.parse_args()
    try:
        if args.operation == "create":
            create(args.scenario, args.workspace)
            report = {"created": True}
        else:
            report = verify(args.scenario, args.workspace)
        print(json.dumps(report))
        sys.exit(0 if report.get("passed", True) else 1)
    except Exception as error:
        print(json.dumps({"passed": False, "errorType": type(error).__name__}))
        sys.exit(1)
